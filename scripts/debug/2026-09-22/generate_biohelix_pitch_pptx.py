#!/usr/bin/env python3
"""
Generate BioHelix Business Plan Roadshow Presentation (PPTX) - Executive Professional Edition
Fully audited: Zero prompt/meta artifacts, clean visual hierarchy, bulleted punchlines.
"""

import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation (16:9 Widescreen)
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette (Deep Biotech Navy + Emerald + Sky Blue)
BG_COLOR = RGBColor(11, 19, 38)        # #0B1326 Deep Biotech Navy
CARD_BG = RGBColor(19, 31, 58)         # #131F3A Slate Biotech Blue
CARD_BORDER = RGBColor(38, 56, 94)     # #26385E Border Navy
ACCENT_GREEN = RGBColor(52, 211, 153)  # #34D399 Life Science Emerald
ACCENT_BLUE = RGBColor(56, 189, 248)   # #38BDF8 Cyan High-Tech
ACCENT_AMBER = RGBColor(251, 191, 36)  # #FBBF24 Warning / Highlight Gold
TEXT_WHITE = RGBColor(248, 250, 252)   # #F8FAFC Pure White
TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8 Subtitle Gray
TEXT_BODY = RGBColor(203, 213, 225)    # #CBD5E1 Body Gray

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

def add_header(slide, title_text, category_text):
    # Category tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.35))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = category_text.upper()
    p_tag.font.size = Pt(11)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_GREEN
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.78), Inches(11.7), Inches(0.65))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_WHITE

def add_footer(slide, current_page, total_pages=11):
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.92), Inches(11.7), Inches(0.35))
    tf = footer_box.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"BioHelix · 生命科学端-超协同移动工作台  |  商业计划书路演版  |  {current_page} / {total_pages}"
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_MUTED

def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

blank_layout = prs.slide_layouts[6]

card_w = Inches(3.75)
card_h = Inches(2.4)
card_gap = Inches(0.24)
start_x = Inches(0.8)

card_w4 = Inches(2.78)
card_gap4 = Inches(0.2)
start_x4 = Inches(0.8)

# ==========================================
# SLIDE 1: 封面页 (Cover)
# ==========================================
slide1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide1)

create_card(slide1, Inches(0.8), Inches(1.1), Inches(11.733), Inches(5.3), bg_color=RGBColor(16, 26, 50), border_color=RGBColor(35, 55, 95))

badge = create_card(slide1, Inches(1.4), Inches(1.6), Inches(4.5), Inches(0.42), bg_color=RGBColor(20, 50, 40), border_color=ACCENT_GREEN)
tf_badge = badge.text_frame
tf_badge.vertical_anchor = MSO_ANCHOR.MIDDLE
p_b = tf_badge.paragraphs[0]
p_b.text = "生命科学与算力基座专项 · 支持私有化与开放算力"
p_b.font.size = Pt(10.5)
p_b.font.bold = True
p_b.font.color.rgb = ACCENT_GREEN
p_b.alignment = PP_ALIGN.CENTER

title_box1 = slide1.shapes.add_textbox(Inches(1.4), Inches(2.3), Inches(10.5), Inches(2.3))
tf1 = title_box1.text_frame
tf1.word_wrap = True
p1 = tf1.paragraphs[0]
p1.text = "BioHelix 移动端生命科学 AI 执行工作台"
p1.font.size = Pt(34)
p1.font.bold = True
p1.font.color.rgb = TEXT_WHITE

p2 = tf1.add_paragraph()
p2.text = "端-边-超协同的开放智能中枢 —— 让超级算力直达实验现场"
p2.font.size = Pt(17)
p2.font.color.rgb = ACCENT_BLUE
p2.space_before = Pt(12)

p3 = tf1.add_paragraph()
p3.text = "跨平台架构 · 本地沙箱隔离 · 混合算力调度（私有模型 / 商业大模型 / 超算中心）"
p3.font.size = Pt(13)
p3.font.color.rgb = TEXT_MUTED
p3.space_before = Pt(14)

meta_box = slide1.shapes.add_textbox(Inches(1.4), Inches(5.3), Inches(10.5), Inches(0.5))
tf_meta = meta_box.text_frame
p_meta = tf_meta.paragraphs[0]
p_meta.text = "汇报主体：BioHelix 团队   |   轮次诉求：天使轮 / 种子轮   |   2026 年 9 月"
p_meta.font.size = Pt(12)
p_meta.font.color.rgb = TEXT_BODY

# ==========================================
# SLIDE 2: 项目定位与核心价值 (Executive Summary)
# 完全去除长文本段落，改为清晰的战略定义与四大支柱卡片
# ==========================================
slide2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide2)
add_header(slide2, "项目概览：定义生命科学移动端 AI 执行工作台", "Executive Summary")
add_footer(slide2, 2)

# Top Banner: One-sentence core value proposition
top_banner = create_card(slide2, Inches(0.8), Inches(1.55), Inches(11.733), Inches(1.35), bg_color=RGBColor(22, 36, 68), border_color=ACCENT_BLUE)
tf_tb = top_banner.text_frame
tf_tb.word_wrap = True
tf_tb.margin_left = Inches(0.3)
tf_tb.margin_top = Inches(0.18)

p_tb1 = tf_tb.paragraphs[0]
p_tb1.text = "核心使命：打通“湿实验台/现场数据”到“宏观算力支撑”的最后一公里"
p_tb1.font.size = Pt(16)
p_tb1.font.bold = True
p_tb1.font.color.rgb = ACCENT_GREEN

p_tb2 = tf_tb.add_paragraph()
p_tb2.text = "在掌上设备打造兼具本地物理沙箱隔离与异构算力调度的执行工作台，使研究人员随时随地完成生信分析与高通量仿真闭环。"
p_tb2.font.size = Pt(12)
p_tb2.font.color.rgb = TEXT_BODY
p_tb2.space_before = Pt(6)

# 4 Core Value Pillars (2x2 Grid)
val_pillars = [
    ("01 移动端微观数据现场闭环", "针对洁净室、生物安全间、野外采样等免电脑场景，在手机端本地沙箱直接解析 PDB/FASTA/SDF 分子数据，即时清洗与质控。"),
    ("02 开放式混合算力弹性调度", "不锁定单一平台：无缝调度机构局域网私有化大模型（vLLM/Ollama）、商业模型 API 及区域超算中心（如长江3D科算中心）。"),
    ("03 涉密数据物理分级隔离", "端侧采用独立 UID 沙箱与 PRoot 容器，核心化合物结构与菌株序列 100% 物理留端，全流程具备确定性防篡改审计。"),
    ("04 深度赋能垂直产业转化", "覆盖创新药物快速初筛、中医药外业道地药材勘测、合成生物发酵中试车间移动巡检，显著缩短产业化落地周期。"),
]

for i, (vp_title, vp_desc) in enumerate(val_pillars):
    col = i % 2
    row = i // 2
    x = Inches(0.8) + col * Inches(5.95)
    y = Inches(3.1) + row * Inches(1.8)
    
    v_card = create_card(slide2, x, y, Inches(5.75), Inches(1.65))
    tf_v = v_card.text_frame
    tf_v.word_wrap = True
    tf_v.margin_left = Inches(0.25)
    tf_v.margin_right = Inches(0.25)
    tf_v.margin_top = Inches(0.2)
    
    p_vt = tf_v.paragraphs[0]
    p_vt.text = vp_title
    p_vt.font.size = Pt(14)
    p_vt.font.bold = True
    p_vt.font.color.rgb = ACCENT_BLUE if i % 2 == 0 else ACCENT_GREEN
    
    p_vd = tf_v.add_paragraph()
    p_vd.text = vp_desc
    p_vd.font.size = Pt(11)
    p_vd.font.color.rgb = TEXT_BODY
    p_vd.space_before = Pt(6)

# ==========================================
# SLIDE 3: 行业痛点与发展机遇
# ==========================================
slide3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide3)
add_header(slide3, "行业痛点：生物医药研发“重计算在后台，真实验在现场”的断层", "Market Problem")
add_footer(slide3, 3)

pain_points = [
    ("痛点一：实验现场缺乏计算工具", "场景断层", "P2/P3 洁净室、生物反应器车间、野外采样区严禁或无法携带笨重电脑，科研人员在实验台前无法即时清洗数据与调用生信软件。"),
    ("痛点二：异构算力孤岛与交互滞后", "调度断层", "私有集群、商业大模型与超算中心彼此割裂。科研人员离开工位即失联，排队状态难掌握，缺乏随时随地可用的统一移动调度中枢。"),
    ("痛点三：涉密数据与弱网外业困境", "安全断层", "先导化合物与基因序列属于核心商业机密，严禁未经脱敏上传公共平台；在野外中草药调查、边远科考弱网环境下，云端工具完全瘫痪。"),
]

for i, (p_title, p_tag, p_desc) in enumerate(pain_points):
    x = start_x + i * (card_w + card_gap)
    p_card = create_card(slide3, x, Inches(1.55), card_w, Inches(3.6))
    tf_p = p_card.text_frame
    tf_p.word_wrap = True
    tf_p.margin_left = Inches(0.25)
    tf_p.margin_right = Inches(0.25)
    tf_p.margin_top = Inches(0.3)
    
    p_tag_p = tf_p.paragraphs[0]
    p_tag_p.text = p_tag.upper()
    p_tag_p.font.size = Pt(11)
    p_tag_p.font.bold = True
    p_tag_p.font.color.rgb = ACCENT_AMBER
    
    p_tit_p = tf_p.add_paragraph()
    p_tit_p.text = p_title
    p_tit_p.font.size = Pt(15)
    p_tit_p.font.bold = True
    p_tit_p.font.color.rgb = TEXT_WHITE
    p_tit_p.space_before = Pt(8)
    
    p_desc_p = tf_p.add_paragraph()
    p_desc_p.text = p_desc
    p_desc_p.font.size = Pt(11.5)
    p_desc_p.font.color.rgb = TEXT_BODY
    p_desc_p.space_before = Pt(12)

opp_card = create_card(slide3, Inches(0.8), Inches(5.35), Inches(11.733), Inches(1.25), bg_color=RGBColor(20, 40, 50), border_color=ACCENT_GREEN)
tf_opp = opp_card.text_frame
tf_opp.word_wrap = True
tf_opp.margin_left = Inches(0.3)
tf_opp.margin_top = Inches(0.2)
p_opp = tf_opp.paragraphs[0]
p_opp.text = "破局解法：BioHelix 契合大赛赛道（三）—— 打造端-边-超协同的生命科学计算基础设施"
p_opp.font.size = Pt(13.5)
p_opp.font.bold = True
p_opp.font.color.rgb = ACCENT_GREEN

p_opp_sub = tf_opp.add_paragraph()
p_opp_sub.text = "打通从“移动端微观实验数据清洗”到“宏观异构算力高通量仿真”的执行链路，将超级算力直达实验第一线！"
p_opp_sub.font.size = Pt(11.5)
p_opp_sub.font.color.rgb = TEXT_BODY
p_opp_sub.space_before = Pt(4)

# ==========================================
# SLIDE 4: 核心产品与解决方案
# ==========================================
slide4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide4)
add_header(slide4, "产品方案：BioHelix —— 面向生命科学的开放式移动 AI 工作台", "Solution")
add_footer(slide4, 4)

pillars = [
    ("01 移动端微观数据本地闭环", "支持系统级目录授权与本地私有存储，内置轻量生信库（Biopython/RDKit），在手机端即可直接解析 PDB/FASTA/SDF 分子文件，完成格式清洗与特征提取。"),
    ("02 开放式混合算力一键调度", "自研 Tool Dispatcher 内置异构算力适配器：自然语言一键下发分子对接、分子动力学模拟或酶催化推演，智能路由至本地私有集群、商业 API 或区域超算中心。"),
    ("03 实时流式终端与长任务对账", "高通量模拟过程通过 PTY 伪终端流式回传手机，实时掌握能量收敛与结合能曲线；断网或切后台后支持持久状态机自动对账，杜绝未知副作用二次误操作。"),
    ("04 涉密新药数据物理分级隔离", "轻量微规则进独立进程沙箱，重度计算进容器化 Linux 运行环境，核心数据物理留端；支持全离线模式与企业私有局域网直连，保障知识产权绝对安全。"),
]

for i, (pil_title, pil_desc) in enumerate(pillars):
    col = i % 2
    row = i // 2
    x = Inches(0.8) + col * Inches(5.95)
    y = Inches(1.55) + row * Inches(2.55)
    
    pil_card = create_card(slide4, x, y, Inches(5.75), Inches(2.35))
    tf_pil = pil_card.text_frame
    tf_pil.word_wrap = True
    tf_pil.margin_left = Inches(0.3)
    tf_pil.margin_right = Inches(0.3)
    tf_pil.margin_top = Inches(0.25)
    
    p_pt = tf_pil.paragraphs[0]
    p_pt.text = pil_title
    p_pt.font.size = Pt(16)
    p_pt.font.bold = True
    p_pt.font.color.rgb = ACCENT_GREEN
    
    p_pd = tf_pil.add_paragraph()
    p_pd.text = pil_desc
    p_pd.font.size = Pt(11.5)
    p_pd.font.color.rgb = TEXT_BODY
    p_pd.space_before = Pt(10)

# ==========================================
# SLIDE 5: 深度技术架构
# ==========================================
slide5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide5)
add_header(slide5, "技术架构：端-超协同、物理隔离沙箱与开放算力抽象", "Technology Architecture")
add_footer(slide5, 5)

# Left Column: Layered Architecture
left_card = create_card(slide5, Inches(0.8), Inches(1.55), Inches(5.75), Inches(5.0))
tf_left = left_card.text_frame
tf_left.word_wrap = True
tf_left.margin_left = Inches(0.3)
tf_left.margin_right = Inches(0.3)
tf_left.margin_top = Inches(0.25)

p_lt = tf_left.paragraphs[0]
p_lt.text = "端-超协同系统架构分层"
p_lt.font.size = Pt(15.5)
p_lt.font.bold = True
p_lt.font.color.rgb = ACCENT_GREEN

arch_layers = [
    ("移动交互层 (Touch UI)", "触控优化界面，支持 3D 分子轻量预览、Diff 比对与实时 Pty 终端输出"),
    ("协同中枢层 (Coordinator)", "统一 Agent 循环，生信上下文组装、多模态输入与任务状态全生命周期管理"),
    ("调度管线层 (Dispatcher)", "Schema 校验 -> 生信权限门禁 -> 策略评估 -> 开放算力 HPC/API 适配下发"),
    ("端侧隔离域 (Edge Runtimes)", "QuickJS 独立 UID 沙箱 (微规则) + PRoot 容器化 Linux (Biopython/RDKit)"),
    ("异构算力池 (Hybrid Compute)", "私有化集群 (vLLM/Slurm) + 商业模型服务 + 公共超算 (如长江3D科算中心)"),
]

for l_name, l_desc in arch_layers:
    p_l1 = tf_left.add_paragraph()
    p_l1.text = f"●  {l_name}"
    p_l1.font.size = Pt(12)
    p_l1.font.bold = True
    p_l1.font.color.rgb = TEXT_WHITE
    p_l1.space_before = Pt(8)
    
    p_l2 = tf_left.add_paragraph()
    p_l2.text = f"    {l_desc}"
    p_l2.font.size = Pt(10.5)
    p_l2.font.color.rgb = TEXT_MUTED

# Right Column: Core Tech Moats
right_card = create_card(slide5, Inches(6.8), Inches(1.55), Inches(5.733), Inches(5.0))
tf_right = right_card.text_frame
tf_right.word_wrap = True
tf_right.margin_left = Inches(0.3)
tf_right.margin_right = Inches(0.3)
tf_right.margin_top = Inches(0.25)

p_rt = tf_right.paragraphs[0]
p_rt.text = "核心技术壁垒与架构创新"
p_rt.font.size = Pt(15.5)
p_rt.font.bold = True
p_rt.font.color.rgb = ACCENT_BLUE

moats = [
    ("免 Root 容器化生信环境移植", "在商用 Android 系统上免 Root 稳定运行 Linux 生信工具链与 C++ 扩展库，兼顾移动安全性、低功耗与专业计算能力。"),
    ("异构算力统一抽象与流式调度", "设计通用 HPC/API 调度协议，无缝适配 Slurm 作业、局域网私有大模型与主流云端 API，实现双向流式通信与长任务监控。"),
    ("确定性审计与高可用对账引擎", "基于嵌入式数据库构建确定性状态机，保障实验操作可追溯不可篡改，异常中断精准对账复原，杜绝模型幻觉导致的数据误删。"),
]

for m_name, m_desc in moats:
    p_m1 = tf_right.add_paragraph()
    p_m1.text = f"★  {m_name}"
    p_m1.font.size = Pt(12.5)
    p_m1.font.bold = True
    p_m1.font.color.rgb = ACCENT_BLUE
    p_m1.space_before = Pt(14)
    
    p_m2 = tf_right.add_paragraph()
    p_m2.text = m_desc
    p_m2.font.size = Pt(11)
    p_m2.font.color.rgb = TEXT_BODY
    p_m2.space_before = Pt(4)

# ==========================================
# SLIDE 6: 科算应用场景 (契合大赛赛道)
# ==========================================
slide6 = prs.slides.add_slide(blank_layout)
set_slide_background(slide6)
add_header(slide6, "科算应用：全方位赋能新药研发、合成生物与中药产业", "Scientific Applications")
add_footer(slide6, 6)

sci_apps = [
    ("场景一：创新药物研发与中药野外勘测", "领域一：新药/中药", "野外道地药材与植物采样无网环境下，手机离线识别物种并进行图谱初筛；联网后联动私有/商业模型或超算集群进行靶点反向筛选。"),
    ("场景二：合成生物中试发酵车间巡检", "领域二：生物制造", "在孝感等生物制造中试车间，巡检人员手持终端读取生物反应器数据，本地评估溶氧与pH；调用酶动力学模型秒级生成补料调控方案。"),
    ("场景三：端-超协同生命计算基础设施", "领域三：主申报赛道", "为高校与药企打造新一代移动生信工作台，实验科学家在洁净室内随时唤醒算力，完成微观参数到宏观高通量仿真的秒级闭环。"),
    ("场景四：智能医疗器械与精准医疗", "领域四：智能器械", "对接便携床旁检验设备（POCT），移动端就地完成原始电生理/光谱信号降噪，联动私有化大模型提供多模态精准辅助诊断。"),
]

for i, (s_title, s_tag, s_desc) in enumerate(sci_apps):
    col = i % 2
    row = i // 2
    x = Inches(0.8) + col * Inches(5.95)
    y = Inches(1.55) + row * Inches(2.55)
    
    s_card = create_card(slide6, x, y, Inches(5.75), Inches(2.35))
    tf_s = s_card.text_frame
    tf_s.word_wrap = True
    tf_s.margin_left = Inches(0.3)
    tf_s.margin_right = Inches(0.3)
    tf_s.margin_top = Inches(0.2)
    
    p_st = tf_s.paragraphs[0]
    p_st.text = s_title
    p_st.font.size = Pt(14.5)
    p_st.font.bold = True
    p_st.font.color.rgb = ACCENT_AMBER
    
    p_stag = tf_s.add_paragraph()
    p_stag.text = f"对接赛道：{s_tag}"
    p_stag.font.size = Pt(11)
    p_stag.font.bold = True
    p_stag.font.color.rgb = ACCENT_GREEN
    p_stag.space_before = Pt(2)
    
    p_sd = tf_s.add_paragraph()
    p_sd.text = s_desc
    p_sd.font.size = Pt(11)
    p_sd.font.color.rgb = TEXT_BODY
    p_sd.space_before = Pt(6)

# ==========================================
# SLIDE 7: 商业模式与产业落地
# ==========================================
slide7 = prs.slides.add_slide(blank_layout)
set_slide_background(slide7)
add_header(slide7, "商业模式：立足科研生态，服务药企研发与生物制造基地", "Business Model")
add_footer(slide7, 7)

biz_models = [
    ("科研/高校免费版", "构建学术生态与用户转化漏斗", [
        "全功能生信文件管理与格式转换",
        "标准微量计算与本地离线沙箱",
        "支持接入用户自有 API (BYOK)",
        "低门槛渗透全国生信高校与课题组",
    ], CARD_BORDER, TEXT_WHITE),
    ("Pro 生信专业版", "面向专业算法工程师与研究员", [
        "容器化生信 Linux 环境持续自愈",
        "多通道异构算力一键调度与管理",
        "实时长任务终端 (PTY) 与无损恢复",
        "支持个人买断与年度维护订阅",
    ], ACCENT_GREEN, ACCENT_GREEN),
    ("基地与企业定制套件", "面向中试车间与涉密药企", [
        "合成生物发酵反应器巡检专属插件",
        "药企涉密局域网私有模型与超算直连",
        "生物传感与工业现场协议软硬件协同",
        "提供机构级合规审计与现场技术保障",
    ], ACCENT_BLUE, ACCENT_BLUE),
]

for i, (bm_title, bm_sub, bm_items, border_c, title_c) in enumerate(biz_models):
    x = start_x + i * (card_w + card_gap)
    b_card = create_card(slide7, x, Inches(1.55), card_w, Inches(5.0), border_color=border_c)
    tf_b = b_card.text_frame
    tf_b.word_wrap = True
    tf_b.margin_left = Inches(0.25)
    tf_b.margin_right = Inches(0.25)
    tf_b.margin_top = Inches(0.3)
    
    p_bt = tf_b.paragraphs[0]
    p_bt.text = bm_title
    p_bt.font.size = Pt(16)
    p_bt.font.bold = True
    p_bt.font.color.rgb = title_c
    
    p_bs = tf_b.add_paragraph()
    p_bs.text = bm_sub
    p_bs.font.size = Pt(11)
    p_bs.font.color.rgb = TEXT_MUTED
    p_bs.space_before = Pt(4)
    
    for item in bm_items:
        p_bi = tf_b.add_paragraph()
        p_bi.text = f"✔  {item}"
        p_bi.font.size = Pt(11.5)
        p_bi.font.color.rgb = TEXT_BODY
        p_bi.space_before = Pt(12)

# ==========================================
# SLIDE 8: 落地规划与生态协同
# ==========================================
slide8 = prs.slides.add_slide(blank_layout)
set_slide_background(slide8)
add_header(slide8, "落地规划：扎实的技术演进与标杆示范基地建设", "Roadmap & Milestones")
add_footer(slide8, 8)

milestones = [
    ("Phase 1: 端侧沙箱就绪", "已交付完成", "● 完成 Android 多执行域与安全沙箱\n● 成功移植移动端 Biopython/RDKit\n● 交付系统级生信文件读写与原子写入\n● 验证离线单机数据清洗闭环"),
    ("Phase 2: 算力调度与对接", "大赛推进期 (1-3月)", "● 示范对接长江3D科学计算中心 API\n● 打通私有化开源模型 (vLLM) 移动端调度\n● 交付流式 PTY 终端与收敛日志推流\n● 开启重点生信实验室小范围定向内测"),
    ("Phase 3: 孝感基地中试试点", "产业试点期 (4-8月)", "● 进驻孝感合成生物制造基地开展试点\n● 部署生物反应器中试掌上巡检方案\n● 联合中医药机构实测野外离线采样\n● 推出首个生命科技商业化订阅包"),
    ("Phase 4: 全行业规模化繁荣", "规模推广期 (9-18月)", "● 接入 100+ 款主流生信/制药工具链\n● 打造生命科学移动算力调度通用标准\n● 覆盖全国 100+ 重点医药高校与药企\n● 实现端-超协同生态规模化盈利"),
]

for i, (m_title, m_time, m_content) in enumerate(milestones):
    x = start_x4 + i * (card_w4 + card_gap4)
    m_card = create_card(slide8, x, Inches(1.55), card_w4, Inches(5.0))
    tf_m = m_card.text_frame
    tf_m.word_wrap = True
    tf_m.margin_left = Inches(0.2)
    tf_m.margin_right = Inches(0.2)
    tf_m.margin_top = Inches(0.3)
    
    p_mt = tf_m.paragraphs[0]
    p_mt.text = m_title
    p_mt.font.size = Pt(13.5)
    p_mt.font.bold = True
    p_mt.font.color.rgb = ACCENT_GREEN if i < 2 else ACCENT_BLUE
    
    p_mtime = tf_m.add_paragraph()
    p_mtime.text = m_time
    p_mtime.font.size = Pt(10.5)
    p_mtime.font.color.rgb = ACCENT_AMBER
    p_mtime.space_before = Pt(4)
    
    p_mc = tf_m.add_paragraph()
    p_mc.text = m_content
    p_mc.font.size = Pt(10.5)
    p_mc.font.color.rgb = TEXT_BODY
    p_mc.space_before = Pt(14)

# ==========================================
# SLIDE 9: 核心团队与竞争优势
# ==========================================
slide9 = prs.slides.add_slide(blank_layout)
set_slide_background(slide9)
add_header(slide9, "核心团队与竞争优势：跨界融合的底层工程与计算生物基因", "Team & Competency")
add_footer(slide9, 9)

team_members = [
    ("移动系统与容器底层专家", "联合创始人 / 系统架构师", "前知名科技企业系统底层工程师，深耕 Linux 容器化、Android 内核及多进程 IPC 通信，主导过大规模移动沙箱与底层虚拟化工程项目。"),
    ("计算生物学与 AI 算法专家", "联合创始人 / 生信科学家", "知名药企与计算生物实验室背景，精通分子对接、高通量分子动力学模拟算法与异构算力集群调度，具备扎实的生信实战经验。"),
    ("生物制造与外业场景专家", "产品与产业化负责人", "具备大型生物化工中试发酵车间工艺经验与重大科研外业科考实操背景，深刻理解湿实验台与生产车间的痛点需求。"),
]

for i, (t_role, t_title, t_desc) in enumerate(team_members):
    x = start_x + i * (card_w + card_gap)
    t_card = create_card(slide9, x, Inches(1.55), card_w, Inches(2.75))
    tf_t = t_card.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = Inches(0.25)
    tf_t.margin_right = Inches(0.25)
    tf_t.margin_top = Inches(0.25)
    
    p_tt = tf_t.paragraphs[0]
    p_tt.text = t_role
    p_tt.font.size = Pt(14.5)
    p_tt.font.bold = True
    p_tt.font.color.rgb = TEXT_WHITE
    
    p_ts = tf_t.add_paragraph()
    p_ts.text = t_title
    p_ts.font.size = Pt(11)
    p_ts.font.color.rgb = ACCENT_GREEN
    p_ts.space_before = Pt(4)
    
    p_td = tf_t.add_paragraph()
    p_td.text = t_desc
    p_td.font.size = Pt(11)
    p_td.font.color.rgb = TEXT_BODY
    p_td.space_before = Pt(8)

moat_bottom = create_card(slide9, Inches(0.8), Inches(4.55), Inches(11.733), Inches(2.0), bg_color=RGBColor(16, 36, 46), border_color=ACCENT_GREEN)
tf_mb = moat_bottom.text_frame
tf_mb.word_wrap = True
tf_mb.margin_left = Inches(0.3)
tf_mb.margin_top = Inches(0.2)

p_mbt = tf_mb.paragraphs[0]
p_mbt.text = "差异化竞争壁垒：端-超协同稀缺性 + 开放中立算力架构"
p_mbt.font.size = Pt(14.5)
p_mbt.font.bold = True
p_mbt.font.color.rgb = ACCENT_GREEN

p_mbb = tf_mb.add_paragraph()
p_mbb.text = (
    "1. 全球稀缺的端-超协同移动基座：打破传统超算必须绑死电脑桌面的宿命，让高通量算力真正渗透进湿实验台与发酵车间；\n"
    "2. 极高工程门槛的免 Root 生信沙箱：在商用 Android 系统下安全调谐容器化生信环境，兼具高安全合规与专业生信计算能力；\n"
    "3. 中立开放的混合算力体系：不绑定特定供应商，全面支持私有化部署大模型、商业 API 与公共超算中心，具备广阔商业空间。"
)
p_mbb.font.size = Pt(11)
p_mbb.font.color.rgb = TEXT_BODY
p_mbb.space_before = Pt(6)

# ==========================================
# SLIDE 10: 融资与产业协同诉求
# ==========================================
slide10 = prs.slides.add_slide(blank_layout)
set_slide_background(slide10)
add_header(slide10, "融资需求与大赛协同：借力科算中心算力，加速孝感中试落地", "Financing & Collaboration")
add_footer(slide10, 10)

card_fin_l = create_card(slide10, Inches(0.8), Inches(1.55), Inches(5.75), Inches(5.0))
tf_fl = card_fin_l.text_frame
tf_fl.word_wrap = True
tf_fl.margin_left = Inches(0.3)
tf_fl.margin_top = Inches(0.3)

p_flt = tf_fl.paragraphs[0]
p_flt.text = "融资诉求与大赛支持"
p_flt.font.size = Pt(15.5)
p_flt.font.bold = True
p_flt.font.color.rgb = ACCENT_GREEN

fin_items = [
    ("融资轮次与诉求", "寻求天使轮 / 种子轮股权融资，出让 10% - 15% 股权"),
    ("长江3D科算中心算力对接", "申请作为首批标杆应用对接 3D 科算中心，打通端-超算力流式通道"),
    ("孝感生物制造基地场景入驻", "申请进驻孝感合成生物制造基地，实地部署发酵罐智能巡检试点"),
    ("55% 资金投向技术研发", "端-超协同调度优化、移动端 3D 分子低功耗流式渲染与生信工具适配"),
]

for fi_title, fi_desc in fin_items:
    p_f1 = tf_fl.add_paragraph()
    p_f1.text = f"◆  {fi_title}"
    p_f1.font.size = Pt(12)
    p_f1.font.bold = True
    p_f1.font.color.rgb = TEXT_WHITE
    p_f1.space_before = Pt(10)
    
    p_f2 = tf_fl.add_paragraph()
    p_f2.text = f"    {fi_desc}"
    p_f2.font.size = Pt(11)
    p_f2.font.color.rgb = TEXT_BODY

card_fin_r = create_card(slide10, Inches(6.8), Inches(1.55), Inches(5.733), Inches(5.0))
tf_fr = card_fin_r.text_frame
tf_fr.word_wrap = True
tf_fr.margin_left = Inches(0.3)
tf_fr.margin_top = Inches(0.3)

p_frt = tf_fr.paragraphs[0]
p_frt.text = "未来 12 个月产业落地目标"
p_frt.font.size = Pt(15.5)
p_frt.font.bold = True
p_frt.font.color.rgb = ACCENT_BLUE

milestone_targets = [
    ("异构算力通道打通", "全面兼容 Slurm/PBS 集群、局域网私有大模型与主流商业 API，形成通用调度标准"),
    ("孝感基地中试示范", "在孝感合成生物制造基地落地示范车间，助力发酵异常预警与中试优化"),
    ("用户与科研合作", "覆盖 30+ 所高校生物医药国家重点实验室，激活 10,000+ 名生信科研人员"),
    ("商业订阅验证", "实现月度经常性收入 (MRR) 突破 50 万元人民币，构建正向商业造血循环"),
]

for mt_title, mt_desc in milestone_targets:
    p_m1 = tf_fr.add_paragraph()
    p_m1.text = f"🎯  {mt_title}"
    p_m1.font.size = Pt(12.5)
    p_m1.font.bold = True
    p_m1.font.color.rgb = ACCENT_BLUE
    p_m1.space_before = Pt(12)
    
    p_m2 = tf_fr.add_paragraph()
    p_m2.text = mt_desc
    p_m2.font.size = Pt(11)
    p_m2.font.color.rgb = TEXT_BODY
    p_m2.space_before = Pt(3)

# ==========================================
# SLIDE 11: 封底
# ==========================================
slide11 = prs.slides.add_slide(blank_layout)
set_slide_background(slide11)

create_card(slide11, Inches(0.8), Inches(1.1), Inches(11.733), Inches(5.3), bg_color=RGBColor(16, 26, 50), border_color=RGBColor(35, 55, 95))

end_box = slide11.shapes.add_textbox(Inches(1.4), Inches(2.2), Inches(10.5), Inches(3.0))
tf_end = end_box.text_frame
tf_end.word_wrap = True

p_e1 = tf_end.paragraphs[0]
p_e1.text = "让超级算力直达生命科学实验第一线"
p_e1.font.size = Pt(32)
p_e1.font.bold = True
p_e1.font.color.rgb = TEXT_WHITE

p_e2 = tf_end.add_paragraph()
p_e2.text = "BioHelix —— 开放连接端侧沙箱、私有化算力与前沿大模型"
p_e2.font.size = Pt(18)
p_e2.font.color.rgb = ACCENT_GREEN
p_e2.space_before = Pt(12)

p_e3 = tf_end.add_paragraph()
p_e3.text = "欢迎各位领导、评委专家指导与交流！\n\n现场系统演示 Demo · 端-超调度实测 · 答辩问答"
p_e3.font.size = Pt(14)
p_e3.font.color.rgb = TEXT_BODY
p_e3.space_before = Pt(20)

# Save presentation
output_path = "/Users/dollars/Helix-biohelix/docs/product/BioHelix_商业计划书_路演版.pptx"
prs.save(output_path)
print(f"BioHelix Roadshow Presentation saved to: {output_path}")
